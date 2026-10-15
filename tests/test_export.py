"""Tests for app.services.export: `build_workbook` and `slugify`."""

from __future__ import annotations

from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.models.category import Category
from app.models.expense import ExpenseRecord, SplitRecord
from app.models.receipt import LineItemRecord
from app.models.trip import Trip
from app.models.user import TripUser
from app.services.export import build_workbook, slugify
from app.services.settlements import compute_balances, settle

CHAT_ID = -100


def _trip(**overrides: object) -> Trip:
    defaults = dict(
        id=1,
        chat_id=CHAT_ID,
        name="Japan Trip!!",
        trip_currency="JPY",
        home_currency="SGD",
        timezone="Asia/Tokyo",
        created_at=datetime.now(UTC),
    )
    defaults.update(overrides)
    return Trip(**defaults)  # type: ignore[arg-type]


def _user(user_id: int, display_name: str) -> TripUser:
    now = datetime.now(UTC)
    return TripUser(
        user_id=user_id,
        username=display_name.lower(),
        display_name=display_name,
        first_seen=now,
        last_seen=now,
        joined=True,
    )


def _expense(
    expense_id: str,
    *,
    message_id: int,
    payer_id: int = 1,
    trip_amount_minor: int = 1200,
    home_amount_minor: int = 1000,
    status: str = "confirmed",
    cancelled_at: datetime | None = None,
) -> ExpenseRecord:
    return ExpenseRecord(
        expense_id=expense_id,
        created_at=datetime(2026, 10, 1, 10, 0, tzinfo=UTC),
        transaction_date=date(2026, 10, 1),
        chat_id=CHAT_ID,
        message_id=message_id,
        created_by_user_id=payer_id,
        payer_user_id=payer_id,
        payer_username="ada",
        payer_display_name="Ada",
        description="Ramen",
        merchant="Ichiran",
        category=Category.FOOD,
        original_amount_minor=trip_amount_minor,
        original_currency="JPY",
        trip_amount_minor=trip_amount_minor,
        home_amount_minor=home_amount_minor,
        rate=Decimal("0.0081"),
        rate_date=date(2026, 9, 30),
        source_type="text",
        split_method="equal",
        status=status,
        cancelled_at=cancelled_at,
        notes=None,
    )


def _split(
    expense_id: str, user_id: int, name: str, trip_minor: int, home_minor: int
) -> SplitRecord:
    return SplitRecord(
        expense_id=expense_id,
        user_id=user_id,
        username=name.lower(),
        display_name=name,
        share_trip_minor=trip_minor,
        share_home_minor=home_minor,
        split_method="equal",
    )


def _build() -> tuple[Trip, bytes]:
    trip = _trip()
    users = [_user(1, "Ada"), _user(2, "Bo")]
    confirmed = _expense("e1", message_id=1, trip_amount_minor=1200, home_amount_minor=1000)
    cancelled = _expense(
        "e2",
        message_id=2,
        trip_amount_minor=600,
        home_amount_minor=500,
        status="cancelled",
        cancelled_at=datetime(2026, 10, 2, tzinfo=UTC),
    )
    expenses = [confirmed, cancelled]
    splits = [
        _split("e1", 1, "Ada", 600, 500),
        _split("e1", 2, "Bo", 600, 500),
        _split("e2", 1, "Ada", 300, 250),
        _split("e2", 2, "Bo", 300, 250),
    ]
    line_items = [
        LineItemRecord(
            expense_id="e1",
            position=0,
            description="Ramen bowl",
            quantity=2,
            amount_trip_minor=1200,
            kind="item",
            allocation_method="everyone",
            assigned_user_ids=[1, 2],
        )
    ]
    workbook = build_workbook(trip, users, expenses, splits, line_items)
    return trip, workbook


def test_sheet_names_and_order() -> None:
    _, workbook_bytes = _build()
    wb = load_workbook(BytesIO(workbook_bytes))
    assert wb.sheetnames == ["Expenses", "Shares", "Line items", "People", "Balances"]


def test_expenses_sheet_headers_and_cancelled_row() -> None:
    _, workbook_bytes = _build()
    wb = load_workbook(BytesIO(workbook_bytes))
    sheet = wb["Expenses"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == [
        "Expense ID",
        "Created At",
        "Transaction Date",
        "Payer Name",
        "Description",
        "Merchant",
        "Category",
        "Original Amount",
        "Original Currency",
        "JPY Amount",
        "SGD Amount",
        "Rate JPY→SGD",
        "Rate Date",
        "Source Type",
        "Split Method",
        "Status",
        "Cancelled At",
        "Notes",
    ]
    assert sheet[1][0].font.bold is True
    assert sheet.freeze_panes == "A2"
    rows = {row[0].value: row for row in sheet.iter_rows(min_row=2)}
    assert rows["e1"][15].value == "confirmed"
    assert rows["e2"][15].value == "cancelled"
    # Amounts land as numbers with the currency's decimal format (spec §10a).
    assert rows["e1"][9].value == 1200  # JPY: 0 decimals, major units == minor units
    assert rows["e1"][9].number_format == "#,##0"
    assert rows["e1"][10].value == 10.0  # SGD: 1000 minor -> 10.00 major
    assert rows["e1"][10].number_format == "#,##0.00"


def test_shares_sheet_matches_splits() -> None:
    _, workbook_bytes = _build()
    wb = load_workbook(BytesIO(workbook_bytes))
    sheet = wb["Shares"]
    assert [cell.value for cell in sheet[1]] == [
        "Expense ID",
        "Participant Name",
        "Share JPY",
        "Share SGD",
        "Split Method",
    ]
    assert sheet.max_row == 5  # header + 4 split rows


def test_line_items_sheet() -> None:
    _, workbook_bytes = _build()
    wb = load_workbook(BytesIO(workbook_bytes))
    sheet = wb["Line items"]
    assert [cell.value for cell in sheet[1]] == [
        "Expense ID",
        "Description",
        "Quantity",
        "Amount JPY",
        "Assigned User IDs",
        "Allocation Method",
        "Kind",
    ]
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert row == ("e1", "Ramen bowl", 2, 1200, "1,2", "everyone", "item")


def test_people_sheet() -> None:
    _, workbook_bytes = _build()
    wb = load_workbook(BytesIO(workbook_bytes))
    sheet = wb["People"]
    assert [cell.value for cell in sheet[1]] == [
        "Telegram User ID",
        "Username",
        "Display Name",
        "First Seen",
        "Last Seen",
        "Joined",
    ]
    assert sheet.max_row == 3  # header + 2 people


def test_balances_sheet_matches_compute_balances_and_settle() -> None:
    trip, workbook_bytes = _build()
    wb = load_workbook(BytesIO(workbook_bytes))
    sheet = wb["Balances"]
    assert [cell.value for cell in sheet[1]] == ["Name", "Paid SGD", "Share SGD", "Net SGD"]

    # Only the confirmed expense (e1) feeds balances; e2 is cancelled.
    confirmed = _expense("e1", message_id=1, trip_amount_minor=1200, home_amount_minor=1000)
    splits = [
        _split("e1", 1, "Ada", 600, 500),
        _split("e1", 2, "Bo", 600, 500),
    ]
    balances = compute_balances([confirmed], splits)
    transfers = settle({uid: b.net_minor for uid, b in balances.items()})

    person_rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] in (None, ""):
            break
        person_rows.append(row)
    assert len(person_rows) == len(balances)
    for row in person_rows:
        name, paid, share, net = row[:4]
        if name == "Ada":
            balance = balances[1]
        else:
            balance = balances[2]
        assert paid == balance.paid_minor / 100
        assert share == balance.share_minor / 100
        assert net == balance.net_minor / 100

    settle_header_row = next(
        i for i, row in enumerate(sheet.iter_rows(values_only=True), start=1) if row[0] == "From"
    )
    settle_rows = list(sheet.iter_rows(min_row=settle_header_row + 1, values_only=True))
    non_empty = [r for r in settle_rows if r[0]]
    assert len(non_empty) == len(transfers)
    # Sprint 9 spec item 9d: compare the actual transfer amounts against
    # settle()'s output, not just the row count.
    names = {1: "Ada", 2: "Bo"}
    actual = [(row[0], row[1], row[2]) for row in non_empty]
    expected = [
        (names[t.from_user_id], names[t.to_user_id], t.amount_minor / 100) for t in transfers
    ]
    assert actual == expected


def test_slugify() -> None:
    assert slugify("Japan Trip!!") == "japan-trip"
    assert slugify("  ") == "trip"
    assert slugify("Kyoto & Osaka 2026") == "kyoto-osaka-2026"
