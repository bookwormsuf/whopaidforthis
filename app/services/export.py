"""Builds the `/export` workbook. Spec references: docs/sprints/09-neon-store-and-
export.md ("Export"), pinned-decisions §10 (column vocabularies carry over), §12
(amount formatting), §18.

Pure function, no I/O, deterministic: `build_workbook` takes the rows a caller
already read from the `LedgerStore` and returns the finished `.xlsx` as bytes, built
in memory with `openpyxl`. The handler (`app.bot.handlers.export`) does the reading
and the `send_document` call.
"""

from __future__ import annotations

import re
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.models.expense import ExpenseRecord, SplitRecord
from app.models.receipt import LineItemRecord
from app.models.trip import Trip
from app.models.user import TripUser
from app.services.currencies import decimals_for
from app.services.settlements import compute_balances, settle

_HEADER_FONT = Font(bold=True)


def _number_format(currency: str) -> str:
    """Spec §10a: `#,##0` for a 0-decimal currency (JPY), else `#,##0.00`."""
    return "#,##0" if decimals_for(currency) == 0 else "#,##0.00"


def _major(amount_minor: int, currency: str) -> float | int:
    """Minor units -> major units as a number, e.g. 1200 JPY -> 1200, 1042 SGD ->
    10.42 -- so the sheet reads naturally (spec: "amounts in major units").
    """
    decimals = decimals_for(currency)
    if decimals == 0:
        return amount_minor
    return float(Decimal(amount_minor) / (10**decimals))


def slugify(text: str) -> str:
    """Lower-case, hyphen-separated slug for a filename. Falls back to "trip" for
    an empty or all-punctuation name.
    """
    slug = re.sub(r"[^a-z0-9]+", "-", text.strip().lower()).strip("-")
    return slug or "trip"


def _write_header(sheet: Worksheet, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
    sheet.freeze_panes = "A2"


def _autosize(sheet: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width


def _name_lookup(
    users: list[TripUser], splits: list[SplitRecord], expenses: list[ExpenseRecord]
) -> dict[int, str]:
    """Freshest name wins: Expenses/Splits rows first, then the Users rows, which
    reflect the person's current name (mirrors `app.bot.handlers.balances._load`).
    """
    names: dict[int, str] = {}
    for split in splits:
        names[split.user_id] = split.display_name
    for expense in expenses:
        names[expense.payer_user_id] = expense.payer_display_name
    for user in users:
        names[user.user_id] = user.display_name
    return names


def _name_for(names: dict[int, str], user_id: int) -> str:
    return names.get(user_id, f"User {user_id}")


def _build_expenses_sheet(wb: Workbook, trip: Trip, expenses: list[ExpenseRecord]) -> None:
    sheet = wb.active
    sheet.title = "Expenses"
    trip_currency, home_currency = trip.trip_currency, trip.home_currency
    _write_header(
        sheet,
        [
            "Expense ID",
            "Created At",
            "Transaction Date",
            "Payer Name",
            "Description",
            "Merchant",
            "Category",
            "Original Amount",
            "Original Currency",
            f"{trip_currency} Amount",
            f"{home_currency} Amount",
            f"Rate {trip_currency}→{home_currency}",
            "Rate Date",
            "Source Type",
            "Split Method",
            "Status",
            "Cancelled At",
            "Notes",
        ],
    )
    trip_fmt = _number_format(trip_currency)
    home_fmt = _number_format(home_currency)
    for expense in expenses:
        original_fmt = _number_format(expense.original_currency)
        row = [
            expense.expense_id,
            expense.created_at.replace(tzinfo=None),
            expense.transaction_date,
            expense.payer_display_name,
            expense.description,
            expense.merchant,
            str(expense.category),
            _major(expense.original_amount_minor, expense.original_currency),
            expense.original_currency,
            _major(expense.trip_amount_minor, trip_currency),
            _major(expense.home_amount_minor, home_currency),
            float(expense.rate),
            expense.rate_date,
            expense.source_type,
            expense.split_method,
            expense.status,
            expense.cancelled_at.replace(tzinfo=None) if expense.cancelled_at else None,
            expense.notes,
        ]
        sheet.append(row)
        last = sheet.max_row
        sheet.cell(row=last, column=3).number_format = "yyyy-mm-dd"
        sheet.cell(row=last, column=8).number_format = original_fmt
        sheet.cell(row=last, column=10).number_format = trip_fmt
        sheet.cell(row=last, column=11).number_format = home_fmt
        sheet.cell(row=last, column=12).number_format = "0.00000000"
        sheet.cell(row=last, column=13).number_format = "yyyy-mm-dd"
    _autosize(sheet, [26, 20, 16, 16, 24, 16, 12, 14, 10, 14, 14, 14, 12, 10, 12, 10, 20, 24])


def _build_shares_sheet(wb: Workbook, trip: Trip, splits: list[SplitRecord]) -> None:
    sheet = wb.create_sheet("Shares")
    trip_currency, home_currency = trip.trip_currency, trip.home_currency
    _write_header(
        sheet,
        [
            "Expense ID",
            "Participant Name",
            f"Share {trip_currency}",
            f"Share {home_currency}",
            "Split Method",
        ],
    )
    trip_fmt = _number_format(trip_currency)
    home_fmt = _number_format(home_currency)
    for split in splits:
        sheet.append(
            [
                split.expense_id,
                split.display_name,
                _major(split.share_trip_minor, trip_currency),
                _major(split.share_home_minor, home_currency),
                split.split_method,
            ]
        )
        last = sheet.max_row
        sheet.cell(row=last, column=3).number_format = trip_fmt
        sheet.cell(row=last, column=4).number_format = home_fmt
    _autosize(sheet, [26, 20, 14, 14, 12])


def _build_line_items_sheet(wb: Workbook, trip: Trip, line_items: list[LineItemRecord]) -> None:
    sheet = wb.create_sheet("Line items")
    trip_currency = trip.trip_currency
    _write_header(
        sheet,
        [
            "Expense ID",
            "Description",
            "Quantity",
            f"Amount {trip_currency}",
            "Assigned User IDs",
            "Allocation Method",
            "Kind",
        ],
    )
    trip_fmt = _number_format(trip_currency)
    for item in line_items:
        sheet.append(
            [
                item.expense_id,
                item.description,
                item.quantity,
                _major(item.amount_trip_minor, trip_currency),
                ",".join(str(uid) for uid in item.assigned_user_ids),
                item.allocation_method,
                item.kind,
            ]
        )
        sheet.cell(row=sheet.max_row, column=4).number_format = trip_fmt
    _autosize(sheet, [26, 24, 10, 14, 24, 16, 12])


def _build_people_sheet(wb: Workbook, users: list[TripUser]) -> None:
    sheet = wb.create_sheet("People")
    _write_header(
        sheet,
        ["Telegram User ID", "Username", "Display Name", "First Seen", "Last Seen", "Joined"],
    )
    for user in users:
        sheet.append(
            [
                user.user_id,
                user.username,
                user.display_name,
                user.first_seen.replace(tzinfo=None),
                user.last_seen.replace(tzinfo=None),
                user.joined,
            ]
        )
    _autosize(sheet, [18, 16, 20, 20, 20, 10])


def _build_balances_sheet(
    wb: Workbook,
    trip: Trip,
    expenses: list[ExpenseRecord],
    splits: list[SplitRecord],
    names: dict[int, str],
) -> None:
    sheet = wb.create_sheet("Balances")
    home_currency = trip.home_currency
    home_fmt = _number_format(home_currency)
    _write_header(
        sheet,
        ["Name", f"Paid {home_currency}", f"Share {home_currency}", f"Net {home_currency}"],
    )

    balances = compute_balances(expenses, splits)
    ordered = sorted(
        balances.values(),
        key=lambda b: (-b.net_minor, _name_for(names, b.user_id).casefold()),
    )
    for balance in ordered:
        sheet.append(
            [
                _name_for(names, balance.user_id),
                _major(balance.paid_minor, home_currency),
                _major(balance.share_minor, home_currency),
                _major(balance.net_minor, home_currency),
            ]
        )
        last = sheet.max_row
        for column in (2, 3, 4):
            sheet.cell(row=last, column=column).number_format = home_fmt

    sheet.append([])  # blank separator row
    header_row = sheet.max_row + 1
    sheet.append(["To settle up", "", "", ""])
    sheet.cell(row=header_row, column=1).font = _HEADER_FONT
    columns_row = sheet.max_row + 1
    sheet.append(["From", "To", f"Amount {home_currency}", ""])
    for cell in sheet[columns_row]:
        cell.font = _HEADER_FONT

    transfers = settle({uid: b.net_minor for uid, b in balances.items()})
    for transfer in transfers:
        sheet.append(
            [
                _name_for(names, transfer.from_user_id),
                _name_for(names, transfer.to_user_id),
                _major(transfer.amount_minor, home_currency),
                "",
            ]
        )
        sheet.cell(row=sheet.max_row, column=3).number_format = home_fmt

    _autosize(sheet, [20, 14, 14, 14])


def build_workbook(
    trip: Trip,
    users: list[TripUser],
    expenses: list[ExpenseRecord],
    splits: list[SplitRecord],
    line_items: list[LineItemRecord],
) -> bytes:
    """One `.xlsx`, five sheets in order: Expenses, Shares, Line items, People,
    Balances. Cancelled expenses are included (with `status`); `Balances` is the
    output of `compute_balances` and `settle`, both confirmed-only already. Pure and
    deterministic -- no model call, nothing written to disk.
    """
    wb = Workbook()
    _build_expenses_sheet(wb, trip, expenses)
    _build_shares_sheet(wb, trip, splits)
    _build_line_items_sheet(wb, trip, line_items)
    _build_people_sheet(wb, users)
    names = _name_lookup(users, splits, expenses)
    _build_balances_sheet(wb, trip, expenses, splits, names)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
